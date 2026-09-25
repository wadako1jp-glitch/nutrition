"""成分表の取り込み・差分・改訂検知ツール（scripts/）のテスト。CI（.github/workflows/ci.yml）で回す。

    python3 -m unittest discover -s tests/python
"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
import diff_food_tables  # noqa: E402
import extract_foods  # noqa: E402
import watch_mext  # noqa: E402

RAW = ROOT / "data" / "mext-tables" / "2023_増補" / "raw" / extract_foods.TABLES["2023_増補"]["xlsx"]
PUBLIC = ROOT / "public" / "data" / extract_foods.TABLES["2023_増補"]["public_name"]


class ExtractTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.current = json.loads(PUBLIC.read_text(encoding="utf-8"))

    def test_配信中のJSONは原本Excelから作り直したものと一致する(self):
        self.assertEqual(extract_foods.extract(RAW), self.current)

    def test_列や行の位置とシート名が変わっても成分識別子で同じ表を読める(self):
        src = openpyxl.load_workbook(RAW, read_only=True, data_only=True)["表全体"]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "全体"
        ws.append(["（改訂で増えた行）"])
        for r in src.iter_rows(values_only=True):
            r = list(r)
            ws.append(r[:5] + ["新しい列"] + r[5:20] + [None, None] + r[20:])
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "shifted.xlsx"
            wb.save(path)
            self.assertEqual(extract_foods.extract(path), self.current)

    def test_成分表の本表でないExcelは読めないと分かる(self):
        wb = openpyxl.Workbook()
        wb.active.append(["アミノ酸成分表"])
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "amino.xlsx"
            wb.save(path)
            with self.assertRaises(extract_foods.LayoutError):
                extract_foods.extract(path)


class DiffTest(unittest.TestCase):
    def test_食品番号の追加と廃止_名前と廃棄率と成分値の変化を報告する(self):
        old = json.loads(PUBLIC.read_text(encoding="utf-8"))
        new = [dict(f) for f in old if f["code"] != "06061"]  # キャベツ（優先表示の番号）がなくなる
        for f in new:
            if f["code"] == "06214":
                f["name"] += "　（改）"
                f["waste_pct"] = 12
        new.append(dict(old[0], code="99001", name="新しい食品"))
        r = diff_food_tables.report(old, new)
        for expected in ("追加された食品番号: **1**", "なくなった食品番号: **1**", "名前が変わった食品: **1**",
                         "廃棄率が変わった食品: **1**", "優先表示（PRIORITY_CODES）でなくなった番号: **1**", "| 06061 |"):
            self.assertIn(expected, r)


class WatchTest(unittest.TestCase):
    def test_ページのリンクから配布ファイルと成分表関連のリンクだけ拾う(self):
        p = watch_mext.LinkParser("https://www.mext.go.jp/a/")
        p.feed('<title> 成分表 </title><a href="x.xlsx">本表</a><a href="/b.html">日本食品標準成分表（九訂）</a>'
               '<a href="/c.html">お知らせ</a>')
        links = [l for l in p.links if watch_mext.interesting(l)]
        self.assertEqual([l["href"] for l in links], ["https://www.mext.go.jp/a/x.xlsx", "https://www.mext.go.jp/b.html"])

    def test_前回からの変化と新しい配布ファイルを見分ける(self):
        old = {"u": {"title": "八訂", "links": [{"href": "https://h/a.xlsx", "text": "本表"}]}}
        new = {"u": {"title": "九訂", "links": [{"href": "https://h/a.xlsx", "text": "本表"},
                                                {"href": "https://h/b.xlsx", "text": "九訂 本表"}]}}
        text, files = watch_mext.diff_snapshots(old, new)
        self.assertEqual(files, ["https://h/b.xlsx"])
        self.assertIn("タイトルが変わった", text)

    def test_今の版のJSONはfoodTable_tsから決まる(self):
        self.assertEqual(watch_mext.current_json(), PUBLIC)


if __name__ == "__main__":
    unittest.main()
